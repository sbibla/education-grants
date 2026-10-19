from pathlib import Path

from docx import Document

from app.db.models import Draft, Grant


def build_docx(draft: Draft, grant: Grant, out_dir: str) -> Path:
    """Write the approved answers to a DOCX the human can send or paste into a portal."""
    doc = Document()
    doc.add_heading(grant.title, level=1)
    if grant.funder:
        doc.add_paragraph(f"Funder: {grant.funder}")
    doc.add_paragraph(f"Source: {grant.url}")
    if grant.fields.get("deadline"):
        doc.add_paragraph(f"Deadline: {grant.fields['deadline']}")
    for f in sorted(draft.fields, key=lambda f: f.id):
        doc.add_heading(f.question, level=2)
        doc.add_paragraph(f.answer)

    path = Path(out_dir) / f"grant-{grant.id}-draft-{draft.id}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
